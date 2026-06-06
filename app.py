import streamlit as st
import google.generativeai as genai
import os
from gtts import gTTS
from docx import Document
from reportlab.pdfgen import canvas
import io

# إعدادات المنصة الأساسية
st.set_page_config(page_title="منصة بصمة التعليمية", layout="wide")

# إخفاء العناصر غير الضرورية
st.markdown("""<style>#MainMenu {visibility: hidden;} footer {visibility: hidden;} header {visibility: hidden;}</style>""", unsafe_allow_html=True)

# تهيئة الذكاء الاصطناعي
if "GOOGLE_API_KEY" in st.secrets:
    genai.configure(api_key=st.secrets["GOOGLE_API_KEY"])

# إدارة الحالة
if 'result_text' not in st.session_state: st.session_state.result_text = ""

# وظائف التصدير (Word, PDF, Audio)
def create_word(text):
    doc = Document()
    doc.add_heading('منصة بصمة للدمج التعليمية', 0)
    doc.add_paragraph(text)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def create_pdf(text):
    bio = io.BytesIO()
    c = canvas.Canvas(bio)
    c.drawString(100, 750, "منصة بصمة للدمج التعليمية")
    c.drawString(100, 730, text[:100])
    c.save()
    return bio.getvalue()

def create_audio(text):
    tts = gTTS(text=text, lang='ar')
    bio = io.BytesIO()
    tts.write_to_fp(bio)
    return bio.getvalue()

# --- الواجهة (التصميم الأصلي) ---
if os.path.exists("logo.jpg"): st.image("logo.jpg", use_column_width=True)
if os.path.exists("waw_logo.png"): st.image("waw_logo.png", width=150)
st.markdown("<h1 style='text-align:center;'>🌟 منصة بصمة للدمج التعليمية 🌟</h1>", unsafe_allow_html=True)

# التبويبات الأصلية
tabs = st.tabs(["💬 الشات الذكي", "🚀 المساعد الذكي", "🎯 أهدافنا", "⚖️ قانون الدمج", "👥 من نحن", "📚 مكتبة المقالات"])

# 1. الشات الذكي
with tabs[0]:
    st.markdown("### 💬 الشات الذكي")
    prompt = st.chat_input("اسألني عن الدمج التعليمي...")
    if prompt:
        model = genai.GenerativeModel('gemini-1.5-flash')
        res = model.generate_content(prompt)
        st.write(res.text)

# 2. المساعد الذكي (التصحيح لـ "حدد" + التحميلات)
with tabs[1]:
    st.markdown("### 🚀 المساعد الذكي")
    c1, c2, c3 = st.columns(3)
    c1.selectbox("حدد المرحلة الدراسية:", ["الأول الابتدائي", "الأول الإعدادي"])
    c2.selectbox("حدد المادة:", ["عربي", "رياضيات", "علوم"])
    c3.selectbox("حدد نوع الإعاقة:", ["توحد", "إعاقة ذهنية"])
    
    st.text_area("وصف النشاط:")
    st.file_uploader("حدد صورة الدرس:")
    if st.button("توليد الخطة"):
        st.session_state.result_text = "هذه خطة درس تفاعلية مخصصة للدمج..."
        st.success("تم توليد الخطة بنجاح!")
    
    if st.session_state.result_text:
        st.write(st.session_state.result_text)
        c_a, c_b, c_c = st.columns(3)
        c_a.download_button("📥 تحميل Word", create_word(st.session_state.result_text), "plan.docx")
        c_b.download_button("📥 تحميل PDF", create_pdf(st.session_state.result_text), "plan.pdf")
        c_c.download_button("📥 تحميل صوت MP3", create_audio(st.session_state.result_text), "plan.mp3")

# 3. الأهداف
with tabs[2]:
    st.markdown("### 🎯 أهدافنا")
    st.write("نحن نسعى لدمج الطلاب ذوي القدرات الخاصة في بيئة تعليمية دامجة ومتميزة.")

# 4. قانون الدمج
with tabs[3]:
    st.markdown("### ⚖️ قانون الدمج")
    st.write("القرار الوزاري 252 لسنة 2017 يضمن حقوق الطلاب المدمجين.")

# 5. من نحن
with tabs[4]:
    st.markdown("### 👥 من نحن")
    st.write("منصة بصمة تهدف لتمكين المعلمين. تم التطوير والبرمجة بواسطة: ولاء مقدام.")
    st.video("https://drive.google.com/file/d/1hGUiJqBkjJhckuO72OMyOul_TtxjTkVJ/preview")

# 6. المقالات
with tabs[5]:
    st.markdown("### 📚 مكتبة المقالات")
    st.write("مقالات حول التوحد، صعوبات التعلم، والتربية الخاصة.")
